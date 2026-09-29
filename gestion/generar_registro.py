"""Genera gestion/plantilla-registro-orbiluz.xlsx: pedidos, clientes, carritos abandonados, devoluciones y resumen.
Uso: python3 gestion/generar_registro.py"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

OUT = os.path.join(os.path.dirname(__file__), "plantilla-registro-orbiluz.xlsx")
N = 300  # filas preparadas
F = "Arial"
NOCHE, AMBAR = "12123A", "FFB547"
H = Font(name=F, bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor=NOCHE)
IN = Font(name=F, color="0000FF"); CALC = Font(name=F, color="000000"); LINK = Font(name=F, color="008000")
AM = PatternFill("solid", fgColor="FFF2CC"); EJ = PatternFill("solid", fgColor="EDEDED")
B = Font(name=F, bold=True); T = Font(name=F, bold=True, size=14, color=NOCHE)
EUR = '#,##0.00 €;-#,##0.00 €;-'; PCT = '0.0%;-0.0%;-'; FECHA = 'dd/mm/yyyy'
thin = Border(bottom=Side(style="thin", color="D9D9D9"))

wb = Workbook()

def hoja(nombre, cols, widths):
    ws = wb.create_sheet(nombre)
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=1, column=i, value=c); cell.font = H; cell.fill = HF
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[cell.column_letter].width = w
    ws.row_dimensions[1].height = 34; ws.freeze_panes = "B2"
    return ws

def lista(ws, rango, opciones):
    dv = DataValidation(type="list", formula1='"' + ",".join(opciones) + '"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(rango)

# ---------- Parámetros ----------
par = wb.active; par.title = "Parámetros"
par["A1"] = "Parámetros (rellena lo amarillo)"; par["A1"].font = T
par["A3"], par["B3"] = "Comisión de pago (%)", 0.019
par["A4"], par["B4"] = "Comisión de pago fija por pedido (€)", 0.25
par["C3"] = "Confírmalo en Shopify › Configuración › Pagos: depende de tu plan y del tipo de tarjeta."
par["A6"] = "Productos"; par["A6"].font = B
for i, h in enumerate(["Producto", "Precio de venta (€)", "Coste proveedor por unidad (€)", "Coste envío proveedor por unidad (€)"]):
    c = par.cell(row=7, column=i + 1, value=h); c.font = H; c.fill = HF; c.alignment = Alignment(wrap_text=True)
productos = [("Lámpara bola de cristal 3D", 24.90), ("Reloj LED 3D de pared y mesa", 29.90),
             ("Proyector de ondas de agua", 29.90), ("Globo terráqueo que levita", 129.00)]
for r, (p, pv) in enumerate(productos, 8):
    par.cell(row=r, column=1, value=p).font = CALC
    par.cell(row=r, column=2, value=pv).font = IN
    for col in (3, 4):
        c = par.cell(row=r, column=col); c.fill = AM; c.font = IN
    for col in (2, 3, 4): par.cell(row=r, column=col).number_format = EUR
notas = ["DSers 29/09/2026: 7,59-7,85 USD + envío 1,99 USD (AliExpress Selection Premium, 5-9 días)",
         "DSers 29/09/2026: 7,51-8,62 USD + envío 1,99 USD (AliExpress Selection Premium, 5-9 días)",
         "DSers 29/09/2026: 11,18 USD + envío 1,99 USD (AliExpress Selection Premium, 5-9 días)",
         "DSers 29/09/2026: 47,43 USD + envío 27,58 USD (1688, envío especial, 10-15 días)"]
for r, n in enumerate(notas, 8): par.cell(row=r, column=5, value=n).font = Font(name=F, italic=True, color="666666")
par["E7"] = "Coste real del proveedor (en USD: pásalo a euros al cambio del día)"; par["E7"].font = H; par["E7"].fill = HF
for col, w in zip("ABCDE", (38, 16, 20, 22, 60)): par.column_dimensions[col].width = w
par["B3"].number_format = PCT; par["B4"].number_format = EUR
for c in ("B3", "B4"): par[c].font = IN; par[c].fill = AM
PROD_RANGE = "Parámetros!$A$8:$A$20"

# ---------- Pedidos ----------
cols = ["Nº pedido", "Fecha", "Nombre cliente", "Email", "Teléfono", "Provincia", "Producto", "Variante", "Unidades",
        "Precio unidad (€)", "Descuento (€)", "Envío cobrado (€)", "Total cobrado (€)", "Coste producto (€)",
        "Coste envío proveedor (€)", "Comisión pago (€)", "Publicidad asignada (€)", "Beneficio (€)", "Margen",
        "Origen", "Estado", "Pedido proveedor (DSers)", "Nº seguimiento", "Fecha envío", "Fecha entrega", "Notas"]
ws = hoja("Pedidos", cols, [11, 12, 20, 26, 14, 14, 30, 16, 9, 12, 11, 11, 13, 13, 13, 12, 13, 12, 9, 13, 13, 18, 20, 12, 12, 30])
for r in range(2, N + 2):
    f = {
        "J": f'=IF(G{r}="","",IFERROR(INDEX(Parámetros!$B$8:$B$20,MATCH(G{r},{PROD_RANGE},0)),""))',
        "M": f'=IF(A{r}="","",I{r}*J{r}-K{r}+L{r})',
        "N": f'=IF(A{r}="","",IFERROR(I{r}*INDEX(Parámetros!$C$8:$C$20,MATCH(G{r},{PROD_RANGE},0)),0))',
        "O": f'=IF(A{r}="","",IFERROR(I{r}*INDEX(Parámetros!$D$8:$D$20,MATCH(G{r},{PROD_RANGE},0)),0))',
        "P": f'=IF(A{r}="","",M{r}*Parámetros!$B$3+Parámetros!$B$4)',
        "R": f'=IF(A{r}="","",M{r}/1.21-N{r}-O{r}-P{r}-Q{r})',
        "S": f'=IF(OR(A{r}="",M{r}=0),"",R{r}/(M{r}/1.21))',
    }
    for col, v in f.items():
        ws[f"{col}{r}"] = v; ws[f"{col}{r}"].font = LINK if col in "JNO" else CALC
    for col in "ABCDEFGHIKLQTUVWXYZ": ws[f"{col}{r}"].font = IN
    for col in "JKLMNOPQR": ws[f"{col}{r}"].number_format = EUR
    ws[f"S{r}"].number_format = PCT
    for col in "BXY": ws[f"{col}{r}"].number_format = FECHA
ws["R1"].comment = Comment("Beneficio sin IVA: Total/1,21 − coste producto − envío proveedor − comisión − publicidad.", "Orbiluz")
ws["J1"].comment = Comment("Se rellena solo desde Parámetros; si vendes a otro precio, escríbelo encima.", "Orbiluz")
lista(ws, f"G2:G{N+1}", [p for p, _ in productos])
lista(ws, f"T2:T{N+1}", ["Meta Ads", "TikTok Ads", "Instagram orgánico", "TikTok orgánico", "Google", "Directo", "Carrito recuperado", "Otro"])
lista(ws, f"U2:U{N+1}", ["Pendiente", "Pedido a proveedor", "Enviado", "Entregado", "Incidencia", "Devuelto", "Cancelado"])

# ---------- Clientes ----------
cols = ["Email", "Nombre", "Teléfono", "Provincia", "Acepta marketing", "Nº pedidos", "Total gastado (€)", "Primer pedido", "Último pedido", "Notas"]
cl = hoja("Clientes", cols, [28, 22, 14, 14, 12, 10, 14, 13, 13, 34])
for r in range(2, N + 2):
    cl[f"F{r}"] = f'=IF(A{r}="","",COUNTIF(Pedidos!$D$2:$D${N+1},A{r}))'
    cl[f"G{r}"] = f'=IF(A{r}="","",SUMIF(Pedidos!$D$2:$D${N+1},A{r},Pedidos!$M$2:$M${N+1}))'
    cl[f"H{r}"] = f'=IF(OR(A{r}="",F{r}=0),"",_xlfn.MINIFS(Pedidos!$B$2:$B${N+1},Pedidos!$D$2:$D${N+1},A{r}))'
    cl[f"I{r}"] = f'=IF(OR(A{r}="",F{r}=0),"",_xlfn.MAXIFS(Pedidos!$B$2:$B${N+1},Pedidos!$D$2:$D${N+1},A{r}))'
    for col in "FGHI": cl[f"{col}{r}"].font = LINK
    for col in "ABCDEJ": cl[f"{col}{r}"].font = IN
    cl[f"G{r}"].number_format = EUR
    for col in "HI": cl[f"{col}{r}"].number_format = FECHA
lista(cl, f"E2:E{N+1}", ["Sí", "No"])

# ---------- Carritos abandonados ----------
cols = ["Fecha", "Email", "Nombre", "Productos", "Importe (€)", "Acepta marketing", "Email 1 (1 h)", "Email 2 (24 h)",
        "Email 3 (48 h)", "Código ofrecido", "Recuperado", "Nº pedido recuperado", "Notas"]
ca = hoja("Carritos abandonados", cols, [12, 28, 20, 34, 12, 12, 12, 12, 12, 14, 11, 14, 30])
for r in range(2, N + 2):
    for col in "ABCDEFGHIJKLM": ca[f"{col}{r}"].font = IN
    ca[f"E{r}"].number_format = EUR
    for col in "AGHI": ca[f"{col}{r}"].number_format = FECHA
lista(ca, f"F2:F{N+1}", ["Sí", "No"]); lista(ca, f"K2:K{N+1}", ["Sí", "No"])
ca["F1"].comment = Comment("Solo puedes enviar emails de recuperación con oferta a quien aceptó marketing o dentro de la venta en curso; revisa la política de privacidad.", "Orbiluz")

# ---------- Devoluciones ----------
cols = ["Nº pedido", "Fecha entrega", "Fecha solicitud", "Días desde entrega", "Dentro de plazo (14 días)", "Motivo", "Estado",
        "Importe reembolsado (€)", "Coste devolución (€)", "Notas"]
dv = hoja("Devoluciones", cols, [11, 13, 13, 11, 12, 30, 14, 14, 14, 30])
for r in range(2, N + 2):
    dv[f"B{r}"] = f'=IF(A{r}="","",IFERROR(INDEX(Pedidos!$Y$2:$Y${N+1},MATCH(A{r},Pedidos!$A$2:$A${N+1},0)),""))'
    dv[f"D{r}"] = f'=IF(OR(B{r}="",C{r}="",B{r}=0),"",C{r}-B{r})'
    dv[f"E{r}"] = f'=IF(D{r}="","",IF(D{r}<=14,"Sí","No"))'
    for col in "BDE": dv[f"{col}{r}"].font = LINK if col == "B" else CALC
    for col in "ACFGHIJ": dv[f"{col}{r}"].font = IN
    for col in "BC": dv[f"{col}{r}"].number_format = FECHA
    for col in "HI": dv[f"{col}{r}"].number_format = EUR
lista(dv, f"F2:F{N+1}", ["Desistimiento (no lo quiere)", "Llegó roto", "No funciona", "No es lo esperado", "Error en el pedido", "No llegó", "Otro"])
lista(dv, f"G2:G{N+1}", ["Solicitada", "Aceptada", "Recibida", "Reembolsada", "Rechazada"])

# ---------- Resumen ----------
rs = wb.create_sheet("Resumen", 0)
rs["A1"] = "Resumen Orbiluz"; rs["A1"].font = T
rs["A2"] = "Se calcula solo. No escribas en esta hoja."
filas = [
    ("Pedidos", '=COUNTA(Pedidos!A2:A%d)' % (N + 1), None),
    ("Pedidos cancelados o devueltos", '=COUNTIF(Pedidos!U2:U%d,"Cancelado")+COUNTIF(Pedidos!U2:U%d,"Devuelto")' % (N + 1, N + 1), None),
    ("Ingresos cobrados con IVA (€)", '=SUM(Pedidos!M2:M%d)' % (N + 1), EUR),
    ("Ingresos sin IVA (€)", '=B6/1.21', EUR),
    ("Beneficio (€)", '=SUM(Pedidos!R2:R%d)' % (N + 1), EUR),
    ("Margen medio", '=IF(B7=0,"",B8/B7)', PCT),
    ("Ticket medio (€)", '=IF(B4=0,"",B6/B4)', EUR),
    ("Publicidad asignada (€)", '=SUM(Pedidos!Q2:Q%d)' % (N + 1), EUR),
    ("ROAS (ingresos / publicidad)", '=IF(B11=0,"",B6/B11)', '0.00'),
    ("Clientes", '=COUNTA(Clientes!A2:A%d)' % (N + 1), None),
    ("Clientes que repiten", '=COUNTIF(Clientes!F2:F%d,">1")' % (N + 1), None),
    ("Carritos abandonados", '=COUNTA(Carritos abandonados!B2:B%d)' % (N + 1), None),
    ("Carritos recuperados", '=COUNTIF(\'Carritos abandonados\'!K2:K%d,"Sí")' % (N + 1), None),
    ("Tasa de recuperación", '=IF(B15=0,"",B16/B15)', PCT),
    ("Devoluciones", '=COUNTA(Devoluciones!A2:A%d)' % (N + 1), None),
    ("Tasa de devolución", '=IF(B4=0,"",B18/B4)', PCT),
]
for i, (lab, fo, fmt) in enumerate(filas, 4):
    rs[f"A{i}"] = lab; rs[f"A{i}"].font = Font(name=F); rs[f"B{i}"] = fo; rs[f"B{i}"].font = B
    if fmt: rs[f"B{i}"].number_format = fmt
    rs[f"A{i}"].border = rs[f"B{i}"].border = thin
rs["B15"] = "=COUNTA('Carritos abandonados'!B2:B%d)" % (N + 1)
rs["D3"] = "Por producto"; rs["D3"].font = B
for j, h in enumerate(["Producto", "Unidades", "Ingresos (€)", "Beneficio (€)"]):
    c = rs.cell(row=4, column=4 + j, value=h); c.font = H; c.fill = HF
for k, (p, _) in enumerate(productos, 5):
    rs[f"D{k}"] = p
    rs[f"E{k}"] = f'=SUMIF(Pedidos!$G$2:$G${N+1},D{k},Pedidos!$I$2:$I${N+1})'
    rs[f"F{k}"] = f'=SUMIF(Pedidos!$G$2:$G${N+1},D{k},Pedidos!$M$2:$M${N+1})'
    rs[f"G{k}"] = f'=SUMIF(Pedidos!$G$2:$G${N+1},D{k},Pedidos!$R$2:$R${N+1})'
    for col in "FG": rs[f"{col}{k}"].number_format = EUR
rs["D11"] = "Por origen"; rs["D11"].font = B
for j, h in enumerate(["Origen", "Pedidos", "Ingresos (€)", "Beneficio (€)"]):
    c = rs.cell(row=12, column=4 + j, value=h); c.font = H; c.fill = HF
for k, o in enumerate(["Meta Ads", "TikTok Ads", "Instagram orgánico", "TikTok orgánico", "Google", "Directo", "Carrito recuperado", "Otro"], 13):
    rs[f"D{k}"] = o
    rs[f"E{k}"] = f'=COUNTIF(Pedidos!$T$2:$T${N+1},D{k})'
    rs[f"F{k}"] = f'=SUMIF(Pedidos!$T$2:$T${N+1},D{k},Pedidos!$M$2:$M${N+1})'
    rs[f"G{k}"] = f'=SUMIF(Pedidos!$T$2:$T${N+1},D{k},Pedidos!$R$2:$R${N+1})'
    for col in "FG": rs[f"{col}{k}"].number_format = EUR
for col, w in zip("ABCDEFG", (32, 16, 3, 30, 11, 14, 14)): rs.column_dimensions[col].width = w
for row in rs.iter_rows():
    for c in row:
        if c.font and c.font.name != F: c.font = Font(name=F, bold=c.font.bold, size=c.font.size, color=c.font.color)

# ---------- Leyenda ----------
lg = wb.create_sheet("Cómo usarlo", 1)
texto = [
    ("Cómo usar este registro", T),
    ("", None),
    ("1. Rellena primero la hoja Parámetros: el coste de cada producto y la comisión de pago (celdas amarillas).", None),
    ("2. Cada venta = una fila en Pedidos (si un pedido lleva 2 productos distintos, usa 2 filas con el mismo Nº de pedido).", None),
    ("3. Escribe solo en las columnas de texto azul. Las negras y verdes se calculan solas.", None),
    ("4. Añade al cliente en Clientes con el mismo email: el nº de pedidos y lo gastado se suman solos.", None),
    ("5. Carritos abandonados: apunta los de Shopify › Pedidos › Carritos abandonados y marca si se recuperan.", None),
    ("6. Devoluciones: el plazo de 14 días se comprueba solo a partir de la fecha de entrega del pedido.", None),
    ("7. La fila 2 de Pedidos es un EJEMPLO: bórrala antes de empezar.", None),
    ("", None),
    ("Protección de datos (RGPD)", B),
    ("Este archivo contiene datos personales. Guarda la copia rellena en gestion/datos/ (está excluida de Git) o en tu ordenador,", None),
    ("nunca en un repositorio ni en una carpeta compartida. Borra los datos de quien lo pida y no los guardes más tiempo del necesario.", None),
    ("", None),
    ("Colores: azul = lo escribes tú · negro = fórmula · verde = viene de otra hoja · amarillo = dato que falta.", None),
]
for i, (t, fnt) in enumerate(texto, 1):
    lg[f"A{i}"] = t; lg[f"A{i}"].font = fnt or Font(name=F)
lg.column_dimensions["A"].width = 120

# Fila de ejemplo (marcada en gris)
ej = {"A": "EJEMPLO-1001", "C": "Cliente Ejemplo", "D": "ejemplo@correo.com", "E": "600000000", "F": "Granada",
      "G": "Lámpara bola de cristal 3D", "H": "Saturno", "I": 2, "K": 9.90, "L": 0, "Q": 6, "T": "Meta Ads", "U": "Enviado",
      "V": "DS-000000", "W": "(nº de seguimiento)"}
import datetime
ws["B2"] = datetime.date(2026, 10, 1); ws["X2"] = datetime.date(2026, 10, 2)
for k, v in ej.items(): ws[f"{k}2"] = v
for c in ws[2]: c.fill = EJ
ws["Z2"] = "Fila de ejemplo: 2 lámparas con la promo 2×39,90 €. Bórrala."

wb.save(OUT); print(OUT)
